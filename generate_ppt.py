import sys
import os

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    COLOR_BG = RGBColor(11, 15, 25)
    COLOR_CARD = RGBColor(22, 28, 45)
    COLOR_PRIMARY = RGBColor(99, 102, 241)
    COLOR_CYAN = RGBColor(6, 182, 212)
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_MUTED = RGBColor(148, 163, 184)

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.color.rgb = COLOR_BG

    def add_header(slide, title_text, category_text="CAMPUSCARE22 PROJECT"):
        add_background(slide)
        
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
        p.font.name = "Arial"

        txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(26)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_WHITE
        p2.font.name = "Arial"

    def add_card(slide, left, top, width, height, title, body_bullets):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_PRIMARY

        txBox = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(width - 0.4), Inches(height - 0.4))
        tf = txBox.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
        p.font.name = "Arial"
        p.space_after = Pt(12)

        for bullet in body_bullets:
            p_bullet = tf.add_paragraph()
            p_bullet.text = "• " + bullet
            p_bullet.font.size = Pt(13)
            p_bullet.font.color.rgb = COLOR_WHITE
            p_bullet.font.name = "Arial"
            p_bullet.space_after = Pt(6)

    # SLIDE 1: Title Slide
    slide1 = prs.slides.add_slide(blank_layout)
    add_background(slide1)

    txBox = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "CAMPUSCARE22"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.font.name = "Arial"

    p2 = tf.add_paragraph()
    p2.text = "College Grievance & Maintenance Issue Tracking Web Platform"
    p2.font.size = Pt(22)
    p2.font.color.rgb = COLOR_WHITE
    p2.font.name = "Arial"
    p2.space_after = Pt(20)

    p3 = tf.add_paragraph()
    p3.text = "Automated Tracking • Role-Based Dashboards • Real-Time Audit Trail"
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_CYAN
    p3.font.name = "Arial"

    # SLIDE 2: Problem & Need
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Problem Statement & Project Need")
    add_card(slide2, 0.8, 1.8, 5.6, 5.0, "Traditional Challenges", [
        "Unorganized verbal/paper complaint filing.",
        "Lack of tracking for students after lodging issues.",
        "Delays in IT, hostel, and electrical maintenance.",
        "No accountability or performance metrics for staff.",
        "Lost ticket records and duplicate reports."
    ])
    add_card(slide2, 6.8, 1.8, 5.6, 5.0, "CampusCare22 Solution", [
        "Digital centralized web portal for all campus issues.",
        "Instant unique Ticket ID generation (e.g. CC22-2026-8941).",
        "Separate Student and Staff authentication portals.",
        "Real-time status timeline (Submitted -> In Progress -> Resolved).",
        "Categorized urgency levels (Low to Emergency)."
    ])

    # SLIDE 3: System Architecture
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "System Architecture & Technology Stack")
    add_card(slide3, 0.8, 1.8, 3.6, 5.0, "Frontend Layer", [
        "HTML5 Semantic Structure",
        "Modern CSS3 Design System",
        "Dark Glassmorphism UI",
        "Vanilla JavaScript (Fetch API)",
        "Responsive Mobile/Desktop Layout"
    ])
    add_card(slide3, 4.8, 1.8, 3.6, 5.0, "Backend Layer (Python)", [
        "Python 3.10+ Runtime",
        "FastAPI Asynchronous Framework",
        "Uvicorn ASGI High Performance Server",
        "RESTful API Routes (Student, Staff, Complaint)",
        "CORS Middleware Enabled"
    ])
    add_card(slide3, 8.8, 1.8, 3.6, 5.0, "Database Layer", [
        "SQLite (Zero-config embedded DB)",
        "MySQL DDL Script (campus_care22.sql)",
        "Normalized Relational Schema",
        "Audit Log History Table",
        "Foreign Key Constraints & Seed Data"
    ])

    # SLIDE 4: Directory Structure
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Project Directory & Code Organization")
    add_card(slide4, 0.8, 1.8, 11.6, 5.0, "CAMPUSCARE22 Repository Structure", [
        "CAMPUSCARE22/frontend/ -> Static Web Pages (index.html, student-login, staff-login, student-dashboard, staff-dashboard, complaint, status)",
        "CAMPUSCARE22/frontend/css/style.css -> Centralized design system with Glassmorphism, animations & CSS variables",
        "CAMPUSCARE22/frontend/js/script.js -> Client API handlers, session state, modal triggers & timeline rendering",
        "CAMPUSCARE22/backend/server.py -> FastAPI backend app mounting routes & serving static files",
        "CAMPUSCARE22/backend/db.py -> SQLite database helper, tables initialization & seed data",
        "CAMPUSCARE22/backend/routes/ -> Modular routers (student.py, staff.py, complaint.py)",
        "CAMPUSCARE22/database/campus_care22.sql -> Complete MySQL relational schema script"
    ])

    # SLIDE 5: Portals
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Core Modules & User Portals")
    add_card(slide5, 0.8, 1.8, 5.6, 5.0, "Student Portal Features", [
        "Secure Roll Number & Password Sign-In.",
        "Lodge complaint with location, category & urgency.",
        "Personal dashboard displaying status counts.",
        "Real-time ticket tracker with step timeline.",
        "View historical complaint log."
    ])
    add_card(slide5, 6.8, 1.8, 5.6, 5.0, "Staff & Admin Control Center", [
        "Staff ID & Employee Credentials Login.",
        "Filter complaints by status (Pending, Review, Resolved).",
        "Filter by department (IT, Hostel, Electrical, Furniture).",
        "Assign tasks to specific maintenance technicians.",
        "Update complaint status & post resolution remarks."
    ])

    # SLIDE 6: Database Schema
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Database Schema & Entity Relationships")
    add_card(slide6, 0.8, 1.8, 5.6, 5.0, "Primary Tables", [
        "students: id, roll_number, name, email, password, department, year_of_study",
        "staff: id, staff_id, name, email, password, department, role",
        "categories: id, name, description, assigned_dept",
        "departments: id, name, code"
    ])
    add_card(slide6, 6.8, 1.8, 5.6, 5.0, "Transactional Tables", [
        "complaints: id, ticket_id, student_id, title, category, location, urgency, description, status, assigned_staff, staff_remarks, created_at",
        "complaint_history: id, complaint_id, status_from, status_to, updated_by_role, updated_by_name, comments, timestamp"
    ])

    # SLIDE 7: Workflow
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Complaint Lifecycle & Workflow")
    add_card(slide7, 0.8, 1.8, 2.7, 5.0, "Step 1: Submission", [
        "Student logs in.",
        "Fills complaint form.",
        "Ticket ID generated (e.g., CC22-2026-8941).",
        "Initial status set to 'Submitted'."
    ])
    add_card(slide7, 3.8, 1.8, 2.7, 5.0, "Step 2: Review", [
        "Staff opens control dashboard.",
        "Reviews complaint urgency & details.",
        "Updates status to 'Under Review'."
    ])
    add_card(slide7, 6.8, 1.8, 2.7, 5.0, "Step 3: In Progress", [
        "Assigns technician/engineer.",
        "Maintenance work initiated.",
        "Status changed to 'In Progress'."
    ])
    add_card(slide7, 9.8, 1.8, 2.7, 5.0, "Step 4: Resolution", [
        "Technician fixes issue.",
        "Staff posts resolution note.",
        "Ticket closed as 'Resolved'.",
        "Timeline updated."
    ])

    # SLIDE 8: Scope & Conclusion
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Future Scope & Conclusion")
    add_card(slide8, 0.8, 1.8, 5.6, 5.0, "Future Scope", [
        "AI-Powered Automatic Categorization.",
        "SMS & Email Notifications.",
        "Photo Attachment Uploads for damaged equipment.",
        "Mobile App Version using React Native / Flutter."
    ])
    add_card(slide8, 6.8, 1.8, 5.6, 5.0, "Conclusion", [
        "CampusCare22 modernizes college grievance handling.",
        "Increases campus maintenance speed and transparency.",
        "Provides robust audit trail and performance analytics.",
        "Fully functional Python backend & modern web frontend."
    ])

    output_path = os.path.join(os.path.dirname(__file__), "CampusCare22_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
